"""将当前SVG转换为可编辑的PowerPoint形状、路径与文字。"""
from pathlib import Path
import re
import math
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont

OUT=Path(__file__).resolve().parent
source=OUT/'问题一_技术路线总图_宋体精简版.svg'
root=ET.parse(source).getroot()
VX,VY,W,H=map(float,root.get('viewBox').split())
SCALE=Inches(20)/W
PT_SCALE=20*72/W
prs=Presentation()
prs.slide_width=round(W*SCALE)
prs.slide_height=round(H*SCALE)
slide=prs.slides.add_slide(prs.slide_layouts[6])
prs.core_properties.title='问题一技术路线图'
prs.core_properties.subject='可编辑矢量流程图'
prs.core_properties.author=''

def node(tag,**attrs):
    e=OxmlElement('a:'+tag)
    for k,v in attrs.items(): e.set(k,str(v))
    return e

def parse(d):
    ts=re.findall(r'[MHVQLZ]|[-+]?(?:\d*\.\d+|\d+)',d)
    i=0; x=y=0; out=[]
    while i<len(ts):
        op=ts[i];i+=1
        if op in ('M','L'):
            x,y=map(float,ts[i:i+2]);i+=2;out.append((op,[(x,y)]))
        elif op=='H':
            x=float(ts[i]);i+=1;out.append(('L',[(x,y)]))
        elif op=='V':
            y=float(ts[i]);i+=1;out.append(('L',[(x,y)]))
        elif op=='Q':
            a,b,x,y=map(float,ts[i:i+4]);i+=4;out.append(('Q',[(a,b),(x,y)]))
        elif op=='Z':out.append(('Z',[]))
        else:raise ValueError(op)
    return out

def rgb(value):
    h=value.lstrip('#')
    return RGBColor.from_string(''.join(c*2 for c in h) if len(h)==3 else h)

def style(shape,fill,stroke,width,dash=False):
    if fill and fill!='none':
        shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    else:shape.fill.background()
    if stroke and stroke!='none':
        shape.line.color.rgb=rgb(stroke)
        shape.line.width=Pt(width*PT_SCALE)
        if dash:
            ln=shape._element.spPr.get_or_add_ln()
            ln.append(node('prstDash',val='dash'))
    else:shape.line.fill.background()

def path_shape(commands,fill='none',stroke=None,width=0,dash=False):
    pts=[p for _,v in commands for p in v]
    x0=min(p[0] for p in pts);y0=min(p[1] for p in pts)
    w=max(max(p[0] for p in pts)-x0,.01)
    h=max(max(p[1] for p in pts)-y0,.01)
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,round((x0-VX)*SCALE),round((y0-VY)*SCALE),max(1,round(w*SCALE)),max(1,round(h*SCALE)))
    sppr=sh._element.spPr
    sppr.remove(sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom'))
    geom=node('custGeom')
    for tag in ['avLst','gdLst','ahLst','cxnLst']:geom.append(node(tag))
    geom.append(node('rect',l='0',t='0',r='r',b='b'))
    paths=node('pathLst');p=node('path',w=max(1,round(w*1000)),h=max(1,round(h*1000)))
    for op,points in commands:
        item=node({'M':'moveTo','L':'lnTo','Q':'quadBezTo','Z':'close'}[op])
        for x,y in points:item.append(node('pt',x=round((x-x0)*1000),y=round((y-y0)*1000)))
        p.append(item)
    paths.append(p);geom.append(paths);sppr.insert(1,geom)
    style(sh,fill,stroke,width,dash)
    return sh

text_records=[]
for element in root:
    tag=element.tag.split('}')[-1]
    if tag=='path':
        d=element.get('d')
        fill=element.get('fill','#20344C');stroke=element.get('stroke')
        width=float(element.get('stroke-width','1'))
        if d=='M0 126H2400V1270H0Z':continue
        for sub in re.findall(r'M[^M]+',d):
            if 'A' in sub:
                nums=re.findall(r'[-+]?(?:\d*\.\d+|\d+)',sub)
                x,y,r=float(nums[0]),float(nums[1]),float(nums[2])
                sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,round((x-2*r-VX)*SCALE),round((y-r-VY)*SCALE),round(2*r*SCALE),round(2*r*SCALE))
                style(sh,fill,stroke,width)
            else:
                commands=parse(sub)
                path_shape(commands,fill,stroke,width,bool(element.get('stroke-dasharray')))
                if element.get('marker-end'):
                    pts=[p for _,v in commands for p in v]
                    a,b=pts[-2:];dx=b[0]-a[0];dy=b[1]-a[1];norm=math.hypot(dx,dy)
                    if norm:
                        ux,uy=dx/norm,dy/norm;length=5*width;half=length/2
                        c=(b[0]-ux*length,b[1]-uy*length)
                        p1=(c[0]-uy*half,c[1]+ux*half);p2=(c[0]+uy*half,c[1]-ux*half)
                        path_shape([('M',[b]),('L',[p1]),('L',[p2]),('Z',[])],'#62758A')
    elif tag=='text':text_records.append(element)

for element in text_records:
    x=float(element.get('x')); y=float(element.get('y')); size=float(element.get('font-size','26'))
    content=''.join(element.itertext());bold=element.get('font-weight')=='bold'
    fontpath='/usr/share/fonts/opentype/noto/NotoSerifCJK-'+('Bold' if bold else 'Regular')+'.ttc'
    font=ImageFont.truetype(fontpath,round(size*10),index=2)
    tw=font.getlength(content)/10+24
    anchor=element.get('text-anchor','middle')
    x0=x-tw/2 if anchor=='middle' else x-12
    sh=slide.shapes.add_textbox(round((x0-VX)*SCALE),round((y-size*.96-VY)*SCALE),round(tw*SCALE),round(size*1.36*SCALE))
    tf=sh.text_frame
    tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    p=tf.paragraphs[0]
    p.alignment=PP_ALIGN.CENTER if anchor=='middle' else PP_ALIGN.LEFT
    p.space_before=p.space_after=Pt(0);p.line_spacing=1.0
    run=p.add_run();run.text=content
    run.font.name='宋体';run.font.size=Pt(size*PT_SCALE);run.font.bold=bold
    run.font.color.rgb=rgb(element.get('fill','#20344C'))
    rpr=run._r.get_or_add_rPr()
    rpr.set('lang','zh-CN');rpr.append(node('ea',typeface='宋体'))
    sh.name=content

path=OUT/'问题一_技术路线总图_可编辑版.pptx'
prs.save(path)
check=Presentation(path)
assert len(check.slides)==1
actual=[sh.text for sh in check.slides[0].shapes if sh.has_text_frame and sh.text]
assert actual==[''.join(e.itertext()) for e in text_records]
assert len({r.font.size for sh in check.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs})==3
print(path)
print(f'1页；{len(actual)}个可编辑文本框；{len(check.slides[0].shapes)}个矢量/文本元素。')
